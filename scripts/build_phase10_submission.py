"""Build final binary submission artifacts from the same evidence used by the gate."""
from __future__ import annotations

import json
import os
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
ART = ROOT / "artifacts"
ART.mkdir(exist_ok=True)


def load(name: str) -> dict:
    return json.loads((ROOT / name).read_text(encoding="utf-8"))


def build_pptx() -> Path:
    from pptx import Presentation
    from pptx.util import Inches, Pt

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slides = [
        ("Swavlamban WAF ML", "Challenge 3 • ML-integrated open-source WAF\nExplainable decisions • controlled learning • real enforcement"),
        ("Security problem", "Adversarial HTTP(S) traffic changes over time. The system combines deterministic rules, adaptive ML signals and human-controlled response without silent model/rule promotion."),
        ("Real request path", "Client → nginx + ModSecurity → Swavlamban WAF gateway → protected upstream\nThe gateway inspects before forwarding and returns 403 without forwarding blocked traffic."),
        ("ML detection", "Supervised + unsupervised anomaly + behavioural burst detector\nVersioned feature schema: http-v2\nEvidence records model/dataset/baseline provenance."),
        ("Explainability", "Every alert can expose detector contributions, active feature groups, rule IDs, reasons, version metadata and a human-readable explanation."),
        ("Rule lifecycle", "Recommend → replay validate → human approve → deploy → enforce → rollback\nApproval identity is tied to the authenticated subject."),
        ("Continuous learning", "Baseline → reviewed feedback → drift → challenger training → evaluation → explicit promotion/rollback\nNo automatic replacement of the active model."),
        ("Operations dashboard", "Authenticated runtime dashboard shows dynamic requests, threats, blocked traffic, model state, telemetry and managed-rule status. Raw request material is not rendered."),
        ("Evidence", f"ModSecurity enforcement: {load('phase10_waf_enforcement_evidence.json').get('sql_blocked_at_waf')} with blocked upstream reach = {load('phase10_waf_enforcement_evidence.json').get('blocked_request_reached_upstream')}\nTLS allow={load('phase10_tls_evidence.json').get('https_allow_status')} • TLS SQL block={load('phase10_tls_evidence.json').get('https_sql_block_status')}\nLoad p95={load('phase10_load_evidence.json')['latency_ms']['p95']} ms"),
        ("Release & boundaries", "Clean-checkout automated gate, full regression, security/RLS verification, binary submission artifacts and portable auditor handoff.\nPublic certificate operations and Internet-scale physical load are intentionally not claimed."),
    ]
    for idx, (title, body) in enumerate(slides, 1):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        bg = slide.background.fill
        bg.solid(); bg.fore_color.rgb = __import__('pptx').dml.color.RGBColor(248,250,252)
        title_box = slide.shapes.add_textbox(Inches(0.7), Inches(0.55), Inches(12), Inches(0.8))
        p = title_box.text_frame.paragraphs[0]
        p.text = f"{idx:02d}  {title}"; p.font.size = Pt(28); p.font.bold = True
        body_box = slide.shapes.add_textbox(Inches(0.8), Inches(1.7), Inches(11.7), Inches(4.8))
        tf = body_box.text_frame; tf.word_wrap = True
        p = tf.paragraphs[0]; p.text = body; p.font.size = Pt(20)
        footer = slide.shapes.add_textbox(Inches(0.8), Inches(6.9), Inches(11.7), Inches(0.25))
        fp = footer.text_frame.paragraphs[0]; fp.text = "swavlamban-waf-ml • Phase 10 audit package"; fp.font.size = Pt(9)
    out = ART / "PHASE10_PRESENTATION.pptx"
    prs.save(out)
    return out


def build_pdf() -> Path:
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak
    from reportlab.lib import colors

    demo = load("phase10_demo_evidence.json")
    load_e = load("phase10_load_evidence.json")
    waf = load("phase10_waf_enforcement_evidence.json")
    tls = load("phase10_tls_evidence.json")
    replay = load("phase10_rule_replay_evidence.json")
    out = ART / "PHASE10_TECHNICAL_REPORT.pdf"
    doc = SimpleDocTemplate(str(out), pagesize=A4, rightMargin=16*mm, leftMargin=16*mm, topMargin=15*mm, bottomMargin=15*mm)
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(name="Small", parent=styles["BodyText"], fontSize=8.5, leading=11))
    story = [Paragraph("Swavlamban WAF ML — Challenge 3 Technical Report", styles["Title"]), Spacer(1, 6)]
    story.append(Paragraph("Architecture and enforcement", styles["Heading2"]))
    story.append(Paragraph("The final candidate combines deterministic WAF signatures, supervised classification, unsupervised anomaly detection and stateful behavioural analysis. The production-oriented request path is client → open-source nginx/ModSecurity enforcement → Swavlamban WAF gateway → protected upstream. Telemetry persistence is isolated from the request decision path through a bounded asynchronous dispatcher.", styles["BodyText"]))
    story.append(Spacer(1, 8))
    story.append(Paragraph("ML and explainability", styles["Heading2"]))
    story.append(Paragraph("Feature schema http-v2 is versioned. Supervised evaluation uses separated train/test partitions; the unsupervised detector trains on a benign-only subset. Decision evidence records detector contributions, feature groups, attributions, reasons, rule IDs and model/dataset/baseline provenance under evidence-v1. Synthetic evaluation is explicitly scoped and is not represented as field accuracy.", styles["BodyText"]))
    story.append(Spacer(1, 8))
    story.append(Paragraph("Rule integration and learning controls", styles["Heading2"]))
    story.append(Paragraph("ML-derived rules are created from decision evidence, replay-validated against positive and negative examples, then require authenticated human approval before deployment. The continuous-learning path uses baseline data, reviewed feedback, drift detection and champion/challenger evaluation with explicit promotion and rollback.", styles["BodyText"]))
    story.append(Spacer(1, 8))
    story.append(Paragraph("Measured evidence", styles["Heading2"]))
    data = [
        ["Check", "Result"],
        ["ModSecurity SQL enforcement", f"blocked={waf.get('sql_blocked_at_waf')} upstream_reached={waf.get('blocked_request_reached_upstream')}"],
        ["TLS termination", f"allow={tls.get('https_allow_status')} SQL block={tls.get('https_sql_block_status')}"],
        ["Rule replay", f"validated={replay['validation']['valid']} positive={replay['positive_example']['matched']} negative={replay['negative_example']['matched']}"],
        ["Demo benchmark", f"500 requests mean={demo['benchmark']['mean_ms']} ms max={demo['benchmark']['max_ms']} ms"],
        ["Network load", f"concurrency={load_e['concurrency']} achieved={load_e['achieved_requests_per_second']} req/s p50={load_e['latency_ms']['p50']} p95={load_e['latency_ms']['p95']} p99={load_e['latency_ms']['p99']} ms"],
    ]
    table = Table(data, colWidths=[65*mm, 105*mm])
    table.setStyle(TableStyle([('GRID',(0,0),(-1,-1),0.4,colors.grey),('BACKGROUND',(0,0),(-1,0),colors.lightgrey),('FONTSIZE',(0,0),(-1,-1),8),('VALIGN',(0,0),(-1,-1),'TOP')]))
    story.append(table)
    story.append(PageBreak())
    story.append(Paragraph("Security and privacy", styles["Heading1"]))
    story.append(Paragraph("Production configuration fails closed when secrets, HTTPS Supabase access, explicit CORS or persistent storage requirements are missing. Bearer tokens are signed with HS256 and checked for issuer, audience, expiry, subject and role. Supabase uses RLS and least-privilege role access. Runtime telemetry excludes raw bodies, headers and query strings and hashes source identity before persistence.", styles["BodyText"]))
    story.append(Spacer(1, 8))
    story.append(Paragraph("Dashboard and operator experience", styles["Heading2"]))
    story.append(Paragraph("The dashboard is served by the API, requires an authenticated token for runtime data, and renders dynamic requests, threats, telemetry state, model state and rule lifecycle data. Controls for recommendation, validation, approval and deployment map to the same authenticated API permissions. DOM rendering uses textContent for runtime fields rather than injecting arbitrary markup.", styles["BodyText"]))
    story.append(Spacer(1, 8))
    story.append(Paragraph("Scale and reliability boundary", styles["Heading2"]))
    story.append(Paragraph("The repository contains a configurable load harness with duration, concurrency, requested rate, payload profiles and p50/p95/p99 reporting. The final CI run reports only the achieved local values. Millions of Internet requests were not fabricated. Fault handling includes rate limits, size limits, upstream timeout/unavailability responses, asynchronous telemetry backpressure and safe shutdown of gateway processes.", styles["BodyText"]))
    story.append(PageBreak())
    story.append(Paragraph("Reproduction and auditor instructions", styles["Heading1"]))
    story.append(Paragraph("Use a clean checkout of the exact final commit and run: `python -m pytest -q`; `python -m compileall -q waf tests scripts`; `python scripts/phase10_master_exam.py`. The GitHub Actions workflow repeats these steps in a fresh runner, installs nginx + the ModSecurity connector, executes process-level WAF enforcement and builds the portable auditor handoff. The handoff manifest distinguishes measured, simulated, design/projection and externally constrained statements.", styles["BodyText"]))
    story.append(Spacer(1, 8))
    story.append(Paragraph("Evidence boundary", styles["Heading2"]))
    story.append(Paragraph("The local certificate is self-signed and exists only for deterministic TLS termination testing. Public certificate issuance/rotation and Internet-scale distributed traffic are outside the free CI environment. Those limits are preserved in the negative-evidence register rather than being turned into false PASS results.", styles["BodyText"]))
    doc.build(story)
    return out


def main() -> int:
    pptx = build_pptx(); pdf = build_pdf()
    print(json.dumps({"pptx": str(pptx), "pdf": str(pdf)}, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
